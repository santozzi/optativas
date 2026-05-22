"""
Exportación a Excel con openpyxl.
Compatible con la lógica del backend Node.js.
"""

import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill('solid', fgColor='3F4F5F')
HEADER_FONT = Font(bold=True, color='FFFFFF', size=11, name='Calibri')
DATA_FONT = Font(size=10, name='Calibri')
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT = Alignment(horizontal='left', vertical='center', wrap_text=True)
thin = Side(style='thin', color='000000')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEADERS = ['Apellido y Nombre', 'LU', 'Documento', 'Inscripción',
           'Código y Materia', 'Genérica Asociada']
COL_WIDTHS = [30, 12, 14, 14, 22, 40]


def export_to_excel(records, filtro_anio=0):
    """Genera el archivo Excel en memoria."""
    wb = Workbook(write_only=False)
    ws = wb.active
    ws.title = 'Inscripciones'

    # Escribir headers
    for col, h in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER

    # Escribir filas de datos
    row_num = 2
    for idx, record in enumerate(records):
        anuales = record._parse_json(record.anualesJson)
        if filtro_anio > 0:
            anuales = [a for a in anuales if a.get('anio') == filtro_anio]

        mats = record._parse_json(record.materiasJson)
        genericas = []
        for a in anuales:
            for g in a.get('generica', []):
                mat = g.get('materia', {})
                genericas.append({
                    'codigo': g.get('codigo', ''),
                    'nombre': f"{g.get('codigo', '')} - {mat.get('codigo', '')} - {mat.get('nombre', '')}",
                })

        count = max(len(mats), len(genericas), 1)

        for i in range(count):
            mat = mats[i] if i < len(mats) else {}
            gen = genericas[i] if i < len(genericas) else {}
            row_data = [
                record.nombre or '',
                record.lu or '',
                record.documento or '',
                record.inscripcion or '',
                f"{mat.get('codigo', '')} - {mat.get('materia', '')}",
                gen.get('nombre', ''),
            ]
            for col, value in enumerate(row_data, start=1):
                cell = ws.cell(row=row_num, column=col, value=value)
                cell.font = DATA_FONT
                cell.border = BORDER
                cell.alignment = CENTER if col != 1 and col != 3 and col != 6 else LEFT

            row_num += 1

        # Combinar celdas de los campos fijos del registro (primera columna a cuarta)
        if count > 1:
            for col in range(1, 5):  # columnas 1-4 = nombre, LU, documento, inscripcion
                start = ws.cell(row=row_num - count, column=col)
                end = ws.cell(row=row_num - 1, column=col)
                ws.merge_cells(
                    start_row=start.row, start_column=start.column,
                    end_row=end.row, end_column=end.column
                )
                # Re-aplicar estilo a celda combinada
                cell = ws.cell(row=start.row, column=col)
                cell.font = DATA_FONT
                cell.border = BORDER
                cell.alignment = LEFT if col in (1, 3) else CENTER

    # Anchos de columna
    for i, width in enumerate(COL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()